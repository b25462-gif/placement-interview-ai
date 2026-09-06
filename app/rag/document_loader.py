"""
Document Loader — loads resumes, JDs, and question bank files into text chunks.
Supports PDF, DOCX, and plain TXT formats.
"""

import os
from pathlib import Path
from typing import List, Dict, Any
from loguru import logger

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter


class DocumentLoader:
    def __init__(self, chunk_size: int = 500, chunk_overlap: int = 50):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            separators=["\n\n", "\n", ".", " ", ""],
        )

    def load_text_file(self, path: Path, metadata: Dict[str, Any] = {}) -> List[Document]:
        """Load a plain .txt file."""
        text = path.read_text(encoding="utf-8", errors="ignore")
        return self._split(text, {"source": str(path), "type": "txt", **metadata})

    def load_pdf(self, path: Path, metadata: Dict[str, Any] = {}) -> List[Document]:
        """Load a PDF file using pypdf."""
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            return self._split(text, {"source": str(path), "type": "pdf", **metadata})
        except Exception as e:
            logger.warning(f"PDF load failed for {path}: {e}")
            return []

    def load_docx(self, path: Path, metadata: Dict[str, Any] = {}) -> List[Document]:
        """Load a .docx file using python-docx."""
        try:
            from docx import Document as DocxDocument
            doc = DocxDocument(str(path))
            text = "\n".join(p.text for p in doc.paragraphs)
            return self._split(text, {"source": str(path), "type": "docx", **metadata})
        except Exception as e:
            logger.warning(f"DOCX load failed for {path}: {e}")
            return []

    def load_directory(self, directory: str, doc_type: str = "general") -> List[Document]:
        """
        Recursively load all supported files from a directory.
        doc_type: 'resume' | 'job_description' | 'question_bank' | 'general'
        """
        dir_path = Path(directory)
        if not dir_path.exists():
            logger.warning(f"Directory not found: {directory}")
            return []

        docs: List[Document] = []
        for file_path in dir_path.rglob("*"):
            if not file_path.is_file():
                continue
            meta = {"doc_type": doc_type, "filename": file_path.name}
            ext = file_path.suffix.lower()
            if ext == ".txt":
                docs.extend(self.load_text_file(file_path, meta))
            elif ext == ".pdf":
                docs.extend(self.load_pdf(file_path, meta))
            elif ext == ".docx":
                docs.extend(self.load_docx(file_path, meta))
            else:
                logger.debug(f"Skipping unsupported file: {file_path.name}")

        logger.info(f"Loaded {len(docs)} chunks from {directory} (type={doc_type})")
        return docs

    def load_all_data(
        self,
        resumes_dir: str,
        jd_dir: str,
        qbank_dir: str,
    ) -> List[Document]:
        """Load all three data sources and return combined document list."""
        docs = []
        docs.extend(self.load_directory(resumes_dir, "resume"))
        docs.extend(self.load_directory(jd_dir, "job_description"))
        docs.extend(self.load_directory(qbank_dir, "question_bank"))
        logger.info(f"Total documents loaded: {len(docs)}")
        return docs

    def _split(self, text: str, metadata: Dict[str, Any]) -> List[Document]:
        if not text.strip():
            return []
        chunks = self.splitter.split_text(text)
        return [Document(page_content=c, metadata=metadata) for c in chunks]
